"""
Dashboard — 3 seviyeli drill-down, st.metric kartları, modern tablo.
"""

import io
import logging
from typing import Dict, List

import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from services import borc_service, cari_service
from services.ilac_service import get_ilac_map
from utils.calculations import hesapla_kalem_tutari, hesapla_iliski_bakiyesi, to_decimal
from utils.formatting import format_bakiye, format_para, format_tarih, format_birim_fiyat, format_miktar
from config import DURUM_AKTIF, DURUM_SILINDI

logger = logging.getLogger(__name__)


# ─── Durum Yönetimi ──────────────────────────────────────────────────────────

def _init_state():
    defaults = {
        "dashboard_seviye": 1,
        "dashboard_cari_id": None,
        "dashboard_karsi_cari_id": None,
        "dashboard_belge_no": None,
        "goster_silinenleri": False,
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def _go1():
    st.session_state.update({"dashboard_seviye": 1, "dashboard_cari_id": None,
                              "dashboard_karsi_cari_id": None, "dashboard_belge_no": None})

def _go2(cid):
    st.session_state.update({"dashboard_seviye": 2, "dashboard_cari_id": cid,
                              "dashboard_karsi_cari_id": None, "dashboard_belge_no": None})

def _go3(cid, kid):
    st.session_state.update({"dashboard_seviye": 3, "dashboard_cari_id": cid,
                              "dashboard_karsi_cari_id": kid, "dashboard_belge_no": None})


# ─── Renk Yardımcıları ───────────────────────────────────────────────────────

def _renk(tutar: int) -> str:
    if tutar > 0: return "#16A34A"
    if tutar < 0: return "#DC2626"
    return "#64748B"

def _html_bakiye(tutar: int, buyuk: bool = False) -> str:
    boyut = "1.1rem" if buyuk else "0.95rem"
    return (f'<span style="color:{_renk(tutar)}; font-weight:700; font-size:{boyut}">'
            f'{format_bakiye(tutar)}</span>')


# ─── Excel Export ─────────────────────────────────────────────────────────────

def _excel_dashboard(data: List[Dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    hf = PatternFill("solid", fgColor="1E293B")
    hfont = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(["Cari Adı", "Net Bakiye (TL)", "Durum"], 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = hf; c.font = hfont
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row["cari_adi"])
        ws.cell(row=i, column=2, value=row["net_bakiye"])
        ws.cell(row=i, column=3, value=row["durum"])
        if row["net_bakiye"] > 0:
            ws.cell(row=i, column=2).font = Font(color="166534", bold=True)
        elif row["net_bakiye"] < 0:
            ws.cell(row=i, column=2).font = Font(color="991B1B", bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 10
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _excel_hareket(hareketler: List[Dict], cari_a_adi: str, cari_b_adi: str,
                   cari_a: int, ilac_map: Dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hareketler"
    hf = PatternFill("solid", fgColor="1E293B")
    hfont = Font(bold=True, color="FFFFFF")
    basliklar = ["Belge No", "Lot Tarihi", "Hareket", "İlaç", "Miktar",
                 "Birim Fiyat", "Tutar (TL)", "Kümülatif (TL)", "Durum"]
    for col, h in enumerate(basliklar, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = hf; c.font = hfont
    for i, h in enumerate(hareketler, 2):
        veren = int(h.get("borc_veren_cari_id", 0))
        hareket = "📤 Verdi" if veren == cari_a else "📥 Aldı"
        tutar = hesapla_kalem_tutari(int(h.get("kalem_miktari", 0)), h.get("birim_alis_fiyati", "0"))
        ws.cell(row=i, column=1, value=h.get("Belge No", ""))
        ws.cell(row=i, column=2, value=format_tarih(str(h.get("Lot Tarihi", ""))))
        ws.cell(row=i, column=3, value=hareket)
        ws.cell(row=i, column=4, value=ilac_map.get(int(h.get("İlaç ID", 0)), "?"))
        ws.cell(row=i, column=5, value=int(h.get("kalem_miktari", 0)))
        ws.cell(row=i, column=6, value=float(to_decimal(h.get("birim_alis_fiyati", "0"))))
        ws.cell(row=i, column=7, value=tutar)
        ws.cell(row=i, column=8, value=h.get("kumulatif_bakiye", 0))
        ws.cell(row=i, column=9, value=h.get("durum", ""))
    widths = [16, 12, 12, 22, 8, 14, 14, 16, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ─── Belge Detay Paneli ───────────────────────────────────────────────────────

def _belge_detay(belge_no: str, cari_map: Dict, ilac_map: Dict, user_id: int):
    from utils.calculations import hesapla_belge_toplami, hesapla_kalem_tutari

    satirlar = borc_service.get_belge_detay(belge_no)
    if not satirlar:
        st.warning("Belge bulunamadı.")
        return

    ilk = satirlar[0]
    veren_id = int(ilk.get("borc_veren_cari_id", 0))
    ilac_id = int(ilk.get("İlaç ID", 0))
    lot = format_tarih(str(ilk.get("Lot Tarihi", "")))
    birim = ilk.get("birim_alis_fiyati", "0")
    toplam_miktar = int(ilk.get("Toplam Miktar", 0))
    durum = str(ilk.get("durum", ""))
    belge_toplami = hesapla_belge_toplami(toplam_miktar, to_decimal(birim))

    # Başlık
    durum_html = ('<span style="background:#DCFCE7;color:#166534;padding:3px 10px;border-radius:20px;font-size:0.8rem;font-weight:600">✓ Aktif</span>'
                  if durum != DURUM_SILINDI else
                  '<span style="background:#FEE2E2;color:#991B1B;padding:3px 10px;border-radius:20px;font-size:0.8rem;font-weight:600">✕ Silindi</span>')

    st.markdown(f"""
    <div style="background:white; border:1px solid #E2E8F0; border-radius:12px; padding:20px; margin:12px 0;">
        <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:16px;">
            <span style="font-size:1rem; font-weight:700; color:#1E293B;">📄 {belge_no}</span>
            {durum_html}
        </div>
        <div style="display:grid; grid-template-columns:1fr 1fr 1fr; gap:12px; margin-bottom:16px;">
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">BORÇ VEREN</div>
                 <div style="font-weight:600;color:#1E293B">{cari_map.get(veren_id,'?')}</div></div>
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">İLAÇ</div>
                 <div style="font-weight:600;color:#1E293B">{ilac_map.get(ilac_id,'?')}</div></div>
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">LOT / ALIŞ TARİHİ</div>
                 <div style="font-weight:600;color:#1E293B">{lot}</div></div>
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">BİRİM ALIŞ FİYATI</div>
                 <div style="font-weight:600;color:#1E293B">{format_birim_fiyat(birim)}</div></div>
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">TOPLAM MİKTAR</div>
                 <div style="font-weight:600;color:#1E293B">{format_miktar(toplam_miktar)} adet</div></div>
            <div><div style="font-size:0.75rem;color:#64748B;font-weight:500">TOPLAM TUTAR</div>
                 <div style="font-weight:700;color:#2563EB;font-size:1.05rem">{format_para(belge_toplami)}</div></div>
        </div>
        <div style="border-top:1px solid #F1F5F9; padding-top:12px;">
            <div style="font-size:0.75rem;color:#64748B;font-weight:500;margin-bottom:8px">DAĞITIM</div>
    """, unsafe_allow_html=True)

    for s in satirlar:
        alan_id = int(s.get("borc_alan_cari_id", 0))
        miktar = int(s.get("kalem_miktari", 0))
        tutar = hesapla_kalem_tutari(miktar, to_decimal(birim))
        s_durum = "✓" if s.get("durum") != DURUM_SILINDI else "✕"
        s_renk = "#166534" if s.get("durum") != DURUM_SILINDI else "#991B1B"
        st.markdown(f"""
        <div style="display:flex; justify-content:space-between; padding:6px 0; border-bottom:1px solid #F8FAFC;">
            <span style="color:{s_renk}; font-weight:500">{s_durum} {cari_map.get(alan_id,'?')}</span>
            <span style="color:#64748B">{format_miktar(miktar)} adet</span>
            <span style="font-weight:600; color:#1E293B">{format_para(tutar)}</span>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("</div></div>", unsafe_allow_html=True)

    # Silme butonu
    if durum != DURUM_SILINDI:
        olusturan = int(ilk.get("Oluşturan User ID", 0))
        if olusturan == user_id:
            col1, col2 = st.columns([1, 4])
            with col1:
                if st.button("🗑️ Belgeyi Sil", key=f"del_{belge_no}", type="secondary"):
                    st.session_state[f"confirm_{belge_no}"] = True

            if st.session_state.get(f"confirm_{belge_no}"):
                st.warning(f"**{belge_no}** silinecek. Bu işlem geri alınamaz!")
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("🗑️ Evet, Sil", key=f"yes_{belge_no}", type="primary"):
                        try:
                            borc_service.delete_borc_belgesi(belge_no, user_id)
                            st.success("Belge silindi.")
                            st.session_state.pop(f"confirm_{belge_no}", None)
                            st.session_state["dashboard_belge_no"] = None
                            st.rerun()
                        except Exception as e:
                            st.error(str(e))
                with c2:
                    if st.button("Vazgeç", key=f"no_{belge_no}"):
                        st.session_state.pop(f"confirm_{belge_no}", None)
                        st.rerun()
        else:
            st.caption("ℹ️ Bu belgeyi yalnızca oluşturan kullanıcı silebilir.")


# ─── Seviye 1 — Tüm Cariler ──────────────────────────────────────────────────

def _seviye1():
    try:
        data = borc_service.get_dashboard_data()
    except Exception as e:
        st.error(f"Veriler yüklenemedi: {e}")
        return

    # Özet metrikler
    aktifler = [d for d in data if d["durum"] == DURUM_AKTIF]
    alacaklilar = [d for d in aktifler if d["net_bakiye"] > 0]
    borcluler = [d for d in aktifler if d["net_bakiye"] < 0]

    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Toplam Cari", len(aktifler))
    with col2:
        st.metric("Net Alacaklı", len(alacaklilar), help="Net bakiyesi pozitif olan cariler")
    with col3:
        st.metric("Net Borçlu", len(borcluler), help="Net bakiyesi negatif olan cariler")

    st.markdown("---")

    col_export, col_info = st.columns([1, 4])
    with col_export:
        excel = _excel_dashboard(data)
        st.download_button("⬇️ Excel", excel, "dashboard.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           key="dl_dash")
    with col_info:
        st.caption("Bir cariye tıklayarak ikili bakiyeleri görüntüleyin.")

    # Tablo başlığı
    hcols = st.columns([4, 3, 2])
    hcols[0].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">CARİ ADI</span>', unsafe_allow_html=True)
    hcols[1].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">NET BAKİYE</span>', unsafe_allow_html=True)
    hcols[2].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">DURUM</span>', unsafe_allow_html=True)
    st.markdown('<hr style="margin:4px 0 8px 0">', unsafe_allow_html=True)

    for row in data:
        cols = st.columns([4, 3, 2])
        with cols[0]:
            if st.button(f"  {row['cari_adi']}", key=f"l1_{row['cari_id']}", use_container_width=True):
                _go2(row["cari_id"])
                st.rerun()
        with cols[1]:
            st.markdown(_html_bakiye(row["net_bakiye"]), unsafe_allow_html=True)
        with cols[2]:
            if row["durum"] == DURUM_AKTIF:
                st.markdown('<span style="background:#DCFCE7;color:#166534;padding:2px 8px;border-radius:12px;font-size:0.8rem;font-weight:600">Aktif</span>', unsafe_allow_html=True)
            else:
                st.markdown('<span style="background:#FEE2E2;color:#991B1B;padding:2px 8px;border-radius:12px;font-size:0.8rem;font-weight:600">Pasif</span>', unsafe_allow_html=True)


# ─── Seviye 2 — Seçilen Cari ─────────────────────────────────────────────────

def _seviye2(cari_id: int):
    cari_map = {int(c["Cari ID"]): str(c["Cari Adı"]) for c in cari_service.get_all_cariler()}
    cari_adi = cari_map.get(cari_id, f"Cari#{cari_id}")

    # Breadcrumb
    bc1, bc2 = st.columns([1, 8])
    with bc1:
        if st.button("← Geri", key="l2_geri"):
            _go1(); st.rerun()

    try:
        dashboard_data = borc_service.get_dashboard_data()
        net = next((d["net_bakiye"] for d in dashboard_data if d["cari_id"] == cari_id), 0)
        iliskiler = borc_service.get_cari_iliskiler(cari_id)
    except Exception as e:
        st.error(f"Veri yüklenemedi: {e}"); return

    # Başlık + metrik
    st.markdown(f"## 🏥 {cari_adi}")
    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Genel Net Bakiye", format_bakiye(net))
    with col2:
        st.metric("İlişkili Cari Sayısı", len(iliskiler))
    with col3:
        alacaklar = [i for i in iliskiler if i["net_bakiye"] > 0]
        st.metric("Net Alacaklı Olduğu", len(alacaklar))

    st.markdown("---")

    if not iliskiler:
        st.info("Bu carinin henüz hiç hareketi bulunmuyor.")
        return

    st.caption("Bir karşı cariye tıklayarak hareket detayını görüntüleyin.")

    hcols = st.columns([4, 3, 2])
    hcols[0].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">KARŞI CARİ</span>', unsafe_allow_html=True)
    hcols[1].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">NET BAKİYE</span>', unsafe_allow_html=True)
    hcols[2].markdown('<span style="font-size:0.8rem;font-weight:600;color:#64748B;text-transform:uppercase">DURUM</span>', unsafe_allow_html=True)
    st.markdown('<hr style="margin:4px 0 8px 0">', unsafe_allow_html=True)

    for row in iliskiler:
        cols = st.columns([4, 3, 2])
        with cols[0]:
            if st.button(f"  {row['cari_adi']}", key=f"l2_{row['cari_id']}", use_container_width=True):
                _go3(cari_id, row["cari_id"]); st.rerun()
        with cols[1]:
            st.markdown(_html_bakiye(row["net_bakiye"]), unsafe_allow_html=True)
        with cols[2]:
            if row["durum"] == DURUM_AKTIF:
                st.markdown('<span style="background:#DCFCE7;color:#166534;padding:2px 8px;border-radius:12px;font-size:0.8rem">Aktif</span>', unsafe_allow_html=True)
            else:
                st.markdown('<span style="background:#FEE2E2;color:#991B1B;padding:2px 8px;border-radius:12px;font-size:0.8rem">Pasif</span>', unsafe_allow_html=True)


# ─── Seviye 3 — İki Cari Arası ───────────────────────────────────────────────

def _seviye3(cari_a: int, cari_b: int, user_id: int):
    from services.borc_service import get_hareketler_normalized
    ilac_map = get_ilac_map()
    cari_map = {int(c["Cari ID"]): str(c["Cari Adı"]) for c in cari_service.get_all_cariler()}
    cari_a_adi = cari_map.get(cari_a, f"Cari#{cari_a}")
    cari_b_adi = cari_map.get(cari_b, f"Cari#{cari_b}")

    bc1, bc2 = st.columns([1, 8])
    with bc1:
        if st.button("← Geri", key="l3_geri"):
            _go2(cari_a); st.rerun()

    goster = st.toggle("Silinenleri Göster", value=st.session_state.get("goster_silinenleri", False))
    st.session_state["goster_silinenleri"] = goster

    try:
        hareketler = borc_service.get_iliski_hareketleri(cari_a, cari_b, goster)
        tum_h = get_hareketler_normalized()
        net = hesapla_iliski_bakiyesi(tum_h, cari_a, cari_b)
    except Exception as e:
        st.error(f"Yüklenemedi: {e}"); return

    # Başlık + metrikler
    st.markdown(f"## {cari_a_adi} ↔ {cari_b_adi}")

    aktif_h = [h for h in hareketler if h.get("durum") != DURUM_SILINDI]

    col1, col2 = st.columns(2)
    with col1:
        st.metric("Net Bakiye", format_bakiye(net))
    with col2:
        st.metric("Aktif İşlem", len(aktif_h))

    st.markdown("---")

    if not hareketler:
        st.info("Bu ikili arasında hareket bulunmuyor.")
        return

    # Excel export
    col_e, col_i = st.columns([1, 5])
    with col_e:
        excel = _excel_hareket(hareketler, cari_a_adi, cari_b_adi, cari_a, ilac_map)
        st.download_button("⬇️ Excel", excel, f"{cari_a_adi}_{cari_b_adi}.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           key="dl_har")

    # Tablo
    hcols = st.columns([2, 2, 2, 3, 2, 2, 2, 2, 1])
    for col, baslik in zip(hcols, ["BELGE NO", "LOT TARİHİ", "HAREKET", "İLAÇ", "MİKTAR", "BİRİM FİYAT", "TUTAR", "KÜMÜLATİF", "DURUM"]):
        col.markdown(f'<span style="font-size:0.75rem;font-weight:600;color:#64748B">{baslik}</span>', unsafe_allow_html=True)
    st.markdown('<hr style="margin:4px 0 8px 0">', unsafe_allow_html=True)

    secili_belge = st.session_state.get("dashboard_belge_no")

    for h in hareketler:
        veren = int(h.get("borc_veren_cari_id", 0))
        hareket_str = "📤 Verdi" if veren == cari_a else "📥 Aldı"
        hareket_renk = "#1D4ED8" if veren == cari_a else "#7C3AED"
        miktar = int(h.get("kalem_miktari", 0))
        birim = h.get("birim_alis_fiyati", "0")
        tutar = hesapla_kalem_tutari(miktar, birim)
        kumul = h.get("kumulatif_bakiye", 0)
        ilac_adi = ilac_map.get(int(h.get("İlaç ID", 0)), "?")
        lot = format_tarih(str(h.get("Lot Tarihi", "")))
        belge_no = str(h.get("Belge No", ""))
        durum = str(h.get("durum", ""))

        cols = st.columns([2, 2, 2, 3, 2, 2, 2, 2, 1])
        with cols[0]:
            if st.button(belge_no, key=f"bn_{belge_no}_{h.get('Kalem No',0)}", use_container_width=True):
                st.session_state["dashboard_belge_no"] = None if secili_belge == belge_no else belge_no
                st.rerun()
        cols[1].markdown(f'<span style="font-size:0.9rem">{lot}</span>', unsafe_allow_html=True)
        cols[2].markdown(f'<span style="color:{hareket_renk};font-weight:600;font-size:0.85rem">{hareket_str}</span>', unsafe_allow_html=True)
        cols[3].markdown(f'<span style="font-size:0.9rem">{ilac_adi}</span>', unsafe_allow_html=True)
        cols[4].markdown(f'<span style="font-size:0.9rem">{format_miktar(miktar)}</span>', unsafe_allow_html=True)
        cols[5].markdown(f'<span style="font-size:0.85rem;color:#64748B">{format_birim_fiyat(birim)}</span>', unsafe_allow_html=True)
        cols[6].markdown(f'<span style="font-weight:600">{format_para(tutar)}</span>', unsafe_allow_html=True)
        cols[7].markdown(_html_bakiye(kumul), unsafe_allow_html=True)
        if durum == DURUM_SILINDI:
            cols[8].markdown('<span style="color:#DC2626;font-size:0.8rem">✕</span>', unsafe_allow_html=True)
        else:
            cols[8].markdown('<span style="color:#16A34A;font-size:0.8rem">✓</span>', unsafe_allow_html=True)

    # Belge detay
    if secili_belge:
        st.markdown("---")
        _belge_detay(secili_belge, cari_map, ilac_map, user_id)


# ─── Ana Render ───────────────────────────────────────────────────────────────

def render(user: Dict) -> None:
    _init_state()
    user_id = int(user.get("User ID", 0))
    seviye = st.session_state.get("dashboard_seviye", 1)

    st.markdown("# 📊 Dashboard")

    if seviye == 1:
        _seviye1()
    elif seviye == 2:
        cid = st.session_state.get("dashboard_cari_id")
        if cid:
            _seviye2(cid)
        else:
            _go1(); st.rerun()
    elif seviye == 3:
        ca = st.session_state.get("dashboard_cari_id")
        cb = st.session_state.get("dashboard_karsi_cari_id")
        if ca and cb:
            _seviye3(ca, cb, user_id)
        else:
            _go1(); st.rerun()
